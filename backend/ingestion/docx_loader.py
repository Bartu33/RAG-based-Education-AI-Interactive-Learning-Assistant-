from pathlib import Path

from docx import Document

from backend.schemas import Segment


def load_docx(path) -> list[Segment]:
    """One segment per heading section; tables are appended as a final section."""
    path = Path(path)
    doc = Document(str(path))
    segments: list[Segment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush():
        text = "\n".join(buffer).strip()
        if text:
            segments.append(Segment(text=text, source_type="docx",
                                    document=path.name, section=heading))
        buffer.clear()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            flush()
            heading = text
        buffer.append(text)
    flush()

    rows: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
    if rows:
        segments.append(Segment(text="\n".join(rows), source_type="docx",
                                document=path.name, section="Tables"))
    return segments