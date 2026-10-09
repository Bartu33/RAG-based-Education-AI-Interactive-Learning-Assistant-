from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.schemas import Segment


def _shape_text(shape) -> list[str]:
    parts: list[str] = []
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            parts.extend(_shape_text(child))
    return parts


def load_pptx(path) -> list[Segment]:
    """One segment per slide (text, tables and speaker notes)."""
    path = Path(path)
    prs = Presentation(str(path))
    segments: list[Segment] = []
    for number, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        text = "\n".join(parts).strip()
        if not text:
            continue
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
        segments.append(Segment(text=text, source_type="pptx", document=path.name,
                                page=number, section=title or None))
    return segments